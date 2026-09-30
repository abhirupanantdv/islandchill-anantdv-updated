from __future__ import annotations

import html
from io import BytesIO

import frappe
from frappe.utils import getdate

from openpyxl import Workbook
from openpyxl.styles import (
	Alignment,
	Border,
	Font,
	PatternFill,
	Side,
)

from islandchill.islandchill.report.user_activity_summary.user_activity_summary import (
	execute,
)


EMAIL_GROUP_NAME = "User Activity Summary"
REPORT_NAME = "User Activity Summary"


def send_daily_user_activity_email(report_date=None):
	if report_date:
		report_date = getdate(report_date)
	else:
		report_date = getdate()

	filters = frappe._dict(
		{
			"from_date": str(report_date),
			"to_date": str(report_date),
			"user": "",
			"action": "",
			"document_type": "",
			"transaction_type": "",
			"status": "",
			"include_logins": 1,
			"enabled_users_only": 1,
		}
	)

	columns, rows, *_rest = execute(filters)

	rows = rows or []
	columns = columns or []

	recipients = get_user_activity_recipients()

	if not recipients:
		frappe.log_error(
			title="User Activity Summary Email",
			message=(
				"No active email addresses found for: "
				+ EMAIL_GROUP_NAME
			),
		)

		return {
			"status": "failed",
			"report_date": str(report_date),
		}

	date_label = report_date.strftime("%d %B %Y")

	xlsx_content = make_user_activity_xlsx(
		columns=columns,
		rows=rows,
		filters=filters,
		date_label=date_label,
	)

	subject = REPORT_NAME + " - " + date_label

	message = build_user_activity_email(
		rows=rows,
		date_label=date_label,
	)

	filename = (
		"User_Activity_Summary_"
		+ report_date.strftime("%Y-%m-%d")
		+ ".xlsx"
	)

	frappe.sendmail(
		recipients=recipients,
		subject=subject,
		message=message,
		attachments=[
			{
				"fname": filename,
				"fcontent": xlsx_content,
			}
		],
		now=True,
	)

	frappe.db.commit()

	return {
		"status": "success",
		"report_date": str(report_date),
	}


def get_user_activity_recipients():
	# 1. Try configured Email Group first
	if frappe.db.exists("Email Group", EMAIL_GROUP_NAME):
		filters = {"email_group": EMAIL_GROUP_NAME}
		member_meta = frappe.get_meta("Email Group Member")
		if member_meta.has_field("unsubscribed"):
			filters["unsubscribed"] = 0
		recipients = frappe.get_all(
			"Email Group Member",
			filters=filters,
			pluck="email",
			limit_page_length=0,
		)
		cleaned = [e.strip() for e in recipients if e and e.strip()]
		if cleaned:
			return list(dict.fromkeys(cleaned))

	# 2. Try legacy group names
	for legacy_group in ("Daily Document Register", "For Courts Daily Document Register"):
		if frappe.db.exists("Email Group", legacy_group):
			recipients = frappe.get_all(
				"Email Group Member",
				filters={"email_group": legacy_group},
				pluck="email",
				limit_page_length=0,
			)
			cleaned = [e.strip() for e in recipients if e and e.strip()]
			if cleaned:
				return list(dict.fromkeys(cleaned))

	# 3. Fallback to active System Managers
	system_managers = frappe.get_all(
		"Has Role",
		filters={"role": "System Manager", "parenttype": "User"},
		pluck="parent",
	)
	if system_managers:
		active_users = frappe.get_all(
			"User",
			filters={"name": ["in", system_managers], "enabled": 1, "user_type": "System User"},
			pluck="email",
		)
		cleaned = [e.strip() for e in active_users if e and e.strip() and "@" in e]
		if cleaned:
			return list(dict.fromkeys(cleaned))

	return []


def make_user_activity_xlsx(
	columns,
	rows,
	filters,
	date_label,
):
	output = BytesIO()

	workbook = Workbook()

	worksheet = workbook.active
	worksheet.title = "User Activity"

	title_fill = PatternFill(
		fill_type="solid",
		fgColor="DDEEFF",
	)

	info_fill = PatternFill(
		fill_type="solid",
		fgColor="EEF6FC",
	)

	header_fill = PatternFill(
		fill_type="solid",
		fgColor="D8EAF8",
	)

	row_fill_1 = PatternFill(
		fill_type="solid",
		fgColor="F7FBFF",
	)

	row_fill_2 = PatternFill(
		fill_type="solid",
		fgColor="EAF2F8",
	)

	border_side = Side(
		style="thin",
		color="B9CBD9",
	)

	cell_border = Border(
		left=border_side,
		right=border_side,
		top=border_side,
		bottom=border_side,
	)

	fieldnames = []
	headers = []

	for column in columns:
		if isinstance(column, dict):
			fieldname = (
				column.get("fieldname")
				or ""
			)

			label = (
				column.get("label")
				or fieldname
				or ""
			)
		else:
			fieldname = str(column)
			label = str(column)

		fieldnames.append(fieldname)
		headers.append(label)

	total_columns = max(
		len(headers),
		2,
	)

	worksheet.merge_cells(
		start_row=1,
		start_column=1,
		end_row=1,
		end_column=total_columns,
	)

	title_cell = worksheet.cell(
		row=1,
		column=1,
		value=REPORT_NAME,
	)

	title_cell.font = Font(
		bold=True,
		size=16,
		color="174D7A",
	)

	title_cell.fill = title_fill

	title_cell.alignment = Alignment(
		horizontal="left",
		vertical="center",
	)

	worksheet.row_dimensions[1].height = 30

	worksheet.cell(
		row=2,
		column=1,
		value="Report Date",
	)

	worksheet.cell(
		row=2,
		column=2,
		value=date_label,
	)

	worksheet.cell(
		row=3,
		column=1,
		value="Include Login / Logout",
	)

	worksheet.cell(
		row=3,
		column=2,
		value="Yes",
	)

	worksheet.cell(
		row=4,
		column=1,
		value="Enabled Users Only",
	)

	worksheet.cell(
		row=4,
		column=2,
		value="Yes",
	)

	for row_number in (2, 3, 4):
		for column_number in (1, 2):
			cell = worksheet.cell(
				row=row_number,
				column=column_number,
			)

			cell.fill = info_fill
			cell.border = cell_border

			cell.alignment = Alignment(
				vertical="center",
			)

		worksheet.cell(
			row=row_number,
			column=1,
		).font = Font(
			bold=True,
			color="35566F",
		)

	header_row = 6

	for column_number, label in enumerate(
		headers,
		start=1,
	):
		cell = worksheet.cell(
			row=header_row,
			column=column_number,
			value=label,
		)

		cell.fill = header_fill

		cell.font = Font(
			bold=True,
			color="123F64",
		)

		cell.alignment = Alignment(
			horizontal="center",
			vertical="center",
			wrap_text=True,
		)

		cell.border = cell_border

	worksheet.row_dimensions[
		header_row
	].height = 32

	first_data_row = header_row + 1

	for excel_row_number, row in enumerate(
		rows,
		start=first_data_row,
	):
		row_index = (
			excel_row_number
			- first_data_row
		)

		if row_index % 2 == 0:
			row_fill = row_fill_1
		else:
			row_fill = row_fill_2

		for column_number, fieldname in enumerate(
			fieldnames,
			start=1,
		):
			value = (
				row.get(fieldname)
				if hasattr(row, "get")
				else ""
			)

			if value is None:
				value = ""

			if isinstance(
				value,
				(
					dict,
					list,
					tuple,
					set,
				),
			):
				value = str(value)

			cell = worksheet.cell(
				row=excel_row_number,
				column=column_number,
				value=value,
			)

			cell.fill = row_fill
			cell.border = cell_border

			cell.alignment = Alignment(
				vertical="top",
				wrap_text=True,
			)

	number_fields = {
		"change_count",
	}

	for column_number, fieldname in enumerate(
		fieldnames,
		start=1,
	):
		if fieldname not in number_fields:
			continue

		for row_number in range(
			first_data_row,
			worksheet.max_row + 1,
		):
			worksheet.cell(
				row=row_number,
				column=column_number,
			).number_format = "#,##0"

	if "amount" in fieldnames:
		amount_column = (
			fieldnames.index("amount")
			+ 1
		)

		for row_number in range(
			first_data_row,
			worksheet.max_row + 1,
		):
			worksheet.cell(
				row=row_number,
				column=amount_column,
			).number_format = "#,##0.00"

	for column_number in range(
		1,
		worksheet.max_column + 1,
	):
		max_length = 0

		for row_number in range(
			1,
			worksheet.max_row + 1,
		):
			value = worksheet.cell(
				row=row_number,
				column=column_number,
			).value

			if value is None:
				continue

			text = str(value)

			length = min(
				len(text),
				50,
			)

			if length > max_length:
				max_length = length

		width = min(
			max(
				max_length + 3,
				11,
			),
			38,
		)

		column_letter = worksheet.cell(
			row=header_row,
			column=column_number,
		).column_letter

		worksheet.column_dimensions[
			column_letter
		].width = width

	for fieldname, width in {
		"action_on": 22,
		"user": 28,
		"full_name": 25,
		"action": 20,
		"document_type": 30,
		"document_name": 26,
		"transaction_type": 25,
		"party": 28,
		"from_status": 22,
		"to_status": 22,
		"current_status": 22,
		"source": 22,
		"source_reference": 22,
		"details": 55,
		"ip_address": 18,
		"last_login": 22,
		"last_active": 22,
	}.items():
		if fieldname in fieldnames:
			column_number = (
				fieldnames.index(fieldname)
				+ 1
			)

			column_letter = worksheet.cell(
				row=header_row,
				column=column_number,
			).column_letter

			worksheet.column_dimensions[
				column_letter
			].width = width

	worksheet.freeze_panes = "A7"
	worksheet.sheet_view.showGridLines = False

	workbook.save(output)

	output.seek(0)

	return output.getvalue()


def format_number(value):
	try:
		return "{:,}".format(
			int(value or 0)
		)
	except Exception:
		return str(value or 0)


def metric_card(
	label,
	value,
	background="#DDEEFF",
	border="#9EC6EB",
	value_color="#185A94",
):
	return (
		"<td width='33.33%' "
		"valign='top' "
		"style='padding:5px;'>"

		"<table role='presentation' "
		"width='100%' "
		"cellpadding='0' "
		"cellspacing='0'>"

		"<tr>"

		"<td style='"
		"background:"
		+ background
		+ ";"
		"border:1px solid "
		+ border
		+ ";"
		"border-radius:8px;"
		"padding:16px 10px;"
		"text-align:center;'>"

		"<div style='"
		"font-family:Arial,sans-serif;"
		"font-size:11px;"
		"font-weight:600;"
		"color:#374151;"
		"margin-bottom:5px;'>"
		+ html.escape(str(label))
		+ "</div>"

		"<div style='"
		"font-family:Arial,sans-serif;"
		"font-size:24px;"
		"line-height:29px;"
		"font-weight:700;"
		"color:"
		+ value_color
		+ ";'>"
		+ html.escape(
			format_number(value)
		)
		+ "</div>"

		"</td>"

		"</tr>"

		"</table>"

		"</td>"
	)


def build_user_activity_email(
	rows,
	date_label,
):
	total_activities = len(rows)

	users = set()

	logins = 0
	logouts = 0
	created = 0
	edited = 0
	updated = 0
	workflow_changes = 0
	status_changes = 0
	submitted = 0
	cancelled = 0
	reopened = 0

	for row in rows:
		if row.get("user"):
			users.add(
				row.get("user")
			)

		action = (
			row.get("action")
			or ""
		)

		if action == "Login":
			logins += 1

		elif action == "Logout":
			logouts += 1

		elif action == "Created":
			created += 1

		elif action == "Edit":
			edited += 1

		elif action == "Updated":
			updated += 1

		elif action == "Workflow Change":
			workflow_changes += 1

		elif action == "Status Change":
			status_changes += 1

		elif action == "Submit":
			submitted += 1

		elif action == "Cancel":
			cancelled += 1

		elif action == "Reopen":
			reopened += 1

	change_events = (
		edited
		+ updated
		+ workflow_changes
		+ status_changes
	)

	cards = (
		"<table role='presentation' "
		"width='100%' "
		"cellpadding='0' "
		"cellspacing='0' "
		"style='margin-bottom:18px;'>"

		"<tr>"

		+ metric_card(
			"Total Activities",
			total_activities,
			background="#D9EAFC",
			border="#8FBCE8",
			value_color="#155A96",
		)

		+ metric_card(
			"Active Users",
			len(users),
			background="#E7DEFC",
			border="#B9A6EE",
			value_color="#613C9C",
		)

		+ metric_card(
			"Logins",
			logins,
			background="#D9EAFC",
			border="#8FBCE8",
			value_color="#155A96",
		)

		+ "</tr>"

		"<tr>"

		+ metric_card(
			"Logouts",
			logouts,
			background="#E7EBEF",
			border="#B8C2CC",
			value_color="#56616B",
		)

		+ metric_card(
			"Documents Created",
			created,
			background="#DDF1E0",
			border="#99D0A3",
			value_color="#247137",
		)

		+ metric_card(
			"Change Events",
			change_events,
			background="#FDEBD5",
			border="#EAB975",
			value_color="#A45D0A",
		)

		+ "</tr>"

		"<tr>"

		+ metric_card(
			"Submitted",
			submitted,
			background="#D8EFEC",
			border="#8CCEC4",
			value_color="#216F62",
		)

		+ metric_card(
			"Cancelled",
			cancelled,
			background="#F9DDE0",
			border="#E8969E",
			value_color="#A5323D",
		)

		+ metric_card(
			"Workflow Changes",
			workflow_changes,
			background="#E5DCFA",
			border="#B6A2E9",
			value_color="#603A99",
		)

		+ "</tr>"

		"<tr>"

		+ metric_card(
			"Status Changes",
			status_changes,
			background="#FCEEC8",
			border="#E4C263",
			value_color="#87600C",
		)

		+ metric_card(
			"Reopened",
			reopened,
			background="#FDEBD5",
			border="#EAB975",
			value_color="#A45D0A",
		)

		+ metric_card(
			"Edits / Updates",
			edited + updated,
			background="#FDEBD5",
			border="#EAB975",
			value_color="#A45D0A",
		)

		+ "</tr>"

		"</table>"
	)

	table_rows = ""

	for index, row in enumerate(rows):
		if index % 2 == 0:
			row_background = "#F7FBFF"
		else:
			row_background = "#EAF2F8"

		document = (
			row.get("document_name")
			or ""
		)

		document_type = (
			row.get("document_type")
			or ""
		)

		table_rows += (
			"<tr style='"
			"background:"
			+ row_background
			+ ";'>"

			"<td style='"
			"padding:7px;"
			"border:1px solid #C3D1DC;"
			"vertical-align:top;'>"
			+ html.escape(
				str(
					row.get("action_on")
					or ""
				)
			)
			+ "</td>"

			"<td style='"
			"padding:7px;"
			"border:1px solid #C3D1DC;"
			"vertical-align:top;'>"
			+ html.escape(
				str(
					row.get("full_name")
					or row.get("user")
					or ""
				)
			)
			+ "</td>"

			"<td style='"
			"padding:7px;"
			"border:1px solid #C3D1DC;"
			"vertical-align:top;"
			"font-weight:600;'>"
			+ html.escape(
				str(
					row.get("action")
					or ""
				)
			)
			+ "</td>"

			"<td style='"
			"padding:7px;"
			"border:1px solid #C3D1DC;"
			"vertical-align:top;'>"
			+ html.escape(
				str(document_type)
			)
			+ "</td>"

			"<td style='"
			"padding:7px;"
			"border:1px solid #C3D1DC;"
			"vertical-align:top;'>"
			+ html.escape(
				str(document)
			)
			+ "</td>"

			"<td style='"
			"padding:7px;"
			"border:1px solid #C3D1DC;"
			"vertical-align:top;'>"
			+ html.escape(
				str(
					row.get(
						"transaction_type"
					)
					or ""
				)
			)
			+ "</td>"

			"<td style='"
			"padding:7px;"
			"border:1px solid #C3D1DC;"
			"vertical-align:top;'>"
			+ html.escape(
				str(
					row.get(
						"current_status"
					)
					or ""
				)
			)
			+ "</td>"

			"<td style='"
			"padding:7px;"
			"border:1px solid #C3D1DC;"
			"vertical-align:top;"
			"word-break:break-word;'>"
			+ html.escape(
				str(
					row.get("details")
					or ""
				)
			)
			+ "</td>"

			"</tr>"
		)

	if not table_rows:
		table_rows = (
			"<tr>"

			"<td colspan='8' "
			"style='"
			"padding:15px;"
			"border:1px solid #C3D1DC;"
			"text-align:center;"
			"background:#F7FBFF;'>"

			"No user activity recorded for this date."

			"</td>"

			"</tr>"
		)

	return (
		"<div style='"
		"font-family:Arial,sans-serif;"
		"font-size:13px;"
		"color:#1F2937;"
		"max-width:1100px;'>"

		"<p>Dear Team,</p>"

		"<p>"
		"Please find the "
		"<b>User Activity Summary</b> "
		"for <b>"
		+ html.escape(date_label)
		+ "</b>."
		"</p>"

		+ cards +

		"<table "
		"width='100%' "
		"cellpadding='0' "
		"cellspacing='0' "
		"style='"
		"width:100%;"
		"border-collapse:collapse;"
		"font-family:Arial,sans-serif;"
		"font-size:11px;'>"

		"<thead>"

		"<tr style='"
		"background:#D8EAF8;"
		"color:#123F64;'>"

		"<th style='"
		"padding:8px 5px;"
		"border:1px solid #AFC4D4;"
		"text-align:center;'>"
		"When"
		"</th>"

		"<th style='"
		"padding:8px 5px;"
		"border:1px solid #AFC4D4;"
		"text-align:center;'>"
		"Employee / User"
		"</th>"

		"<th style='"
		"padding:8px 5px;"
		"border:1px solid #AFC4D4;"
		"text-align:center;'>"
		"Action"
		"</th>"

		"<th style='"
		"padding:8px 5px;"
		"border:1px solid #AFC4D4;"
		"text-align:center;'>"
		"Document Type"
		"</th>"

		"<th style='"
		"padding:8px 5px;"
		"border:1px solid #AFC4D4;"
		"text-align:center;'>"
		"Document"
		"</th>"

		"<th style='"
		"padding:8px 5px;"
		"border:1px solid #AFC4D4;"
		"text-align:center;'>"
		"Transaction Type"
		"</th>"

		"<th style='"
		"padding:8px 5px;"
		"border:1px solid #AFC4D4;"
		"text-align:center;'>"
		"Current Status"
		"</th>"

		"<th style='"
		"padding:8px 5px;"
		"border:1px solid #AFC4D4;"
		"text-align:center;'>"
		"Activity Detail"
		"</th>"

		"</tr>"

		"</thead>"

		"<tbody>"
		+ table_rows
		+ "</tbody>"

		"</table>"

		"<p style='margin-top:18px;'>"
		"The complete User Activity Summary "
		"is attached in Excel format."
		"</p>"

		"<p>"
		"Regards,<br>"
		"<b>ERPNext System</b>"
		"</p>"

		"</div>"
	)