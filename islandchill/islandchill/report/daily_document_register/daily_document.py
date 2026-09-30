from __future__ import annotations

import html
from io import BytesIO

import frappe
from frappe.utils import getdate

from openpyxl import Workbook
from openpyxl.styles import (
	Alignment,
	Border,
	Font,
	PatternFill,
	Side,
)

from islandchill.islandchill.report.daily_document_register.daily_document_register import (
	execute,
)


EMAIL_GROUP_NAME = "Daily Document Register"
REPORT_NAME = "Daily Document Register"


def send_daily_document_register_email(report_date=None):
	if report_date:
		report_date = getdate(report_date)
	else:
		report_date = getdate()

	filters = frappe._dict(
		{
			"report_view": "Daily Consolidated Summary",
			"from_date": str(report_date),
			"to_date": str(report_date),
			"document_type": "",
			"created_by": "",
			"action": "",
			"transaction_type": "",
			"status": "",
			"include_cancelled": 0,
		}
	)

	columns, rows, *_rest = execute(filters)

	rows = rows or []
	columns = columns or []

	recipients = get_daily_document_register_recipients()

	if not recipients:
		frappe.log_error(
			title="Daily Document Register Email",
			message=(
				"No active email addresses found for: "
				+ EMAIL_GROUP_NAME
			),
		)

		return {
			"status": "failed",
			"report_date": str(report_date),
		}

	date_label = report_date.strftime("%d %B %Y")

	xlsx_content = make_daily_document_register_xlsx(
		columns=columns,
		rows=rows,
		filters=filters,
		date_label=date_label,
	)

	subject = REPORT_NAME + " - " + date_label

	message = build_daily_document_register_email(
		rows=rows,
		date_label=date_label,
	)

	filename = (
		"Daily_Document_Register_"
		+ report_date.strftime("%Y-%m-%d")
		+ ".xlsx"
	)

	frappe.sendmail(
		recipients=recipients,
		subject=subject,
		message=message,
		attachments=[
			{
				"fname": filename,
				"fcontent": xlsx_content,
			}
		],
	)

	return {
		"status": "success",
		"report_date": str(report_date),
	}


def get_daily_document_register_recipients():
	# 1. Try configured Email Group first
	if frappe.db.exists("Email Group", EMAIL_GROUP_NAME):
		filters = {"email_group": EMAIL_GROUP_NAME}
		member_meta = frappe.get_meta("Email Group Member")
		if member_meta.has_field("unsubscribed"):
			filters["unsubscribed"] = 0
		recipients = frappe.get_all(
			"Email Group Member",
			filters=filters,
			pluck="email",
			limit_page_length=0,
		)
		cleaned = [e.strip() for e in recipients if e and e.strip()]
		if cleaned:
			return list(dict.fromkeys(cleaned))

	# 2. Try "For Courts Daily Document Register" for legacy compatibility
	if frappe.db.exists("Email Group", "For Courts Daily Document Register"):
		recipients = frappe.get_all(
			"Email Group Member",
			filters={"email_group": "For Courts Daily Document Register"},
			pluck="email",
			limit_page_length=0,
		)
		cleaned = [e.strip() for e in recipients if e and e.strip()]
		if cleaned:
			return list(dict.fromkeys(cleaned))

	# 3. Fallback to active System Managers so email is never lost
	system_managers = frappe.get_all(
		"Has Role",
		filters={"role": "System Manager", "parenttype": "User"},
		pluck="parent",
	)
	if system_managers:
		active_users = frappe.get_all(
			"User",
			filters={"name": ["in", system_managers], "enabled": 1, "user_type": "System User"},
			pluck="email",
		)
		cleaned = [e.strip() for e in active_users if e and e.strip() and "@" in e]
		if cleaned:
			return list(dict.fromkeys(cleaned))

	return []


def make_daily_document_register_xlsx(
	columns,
	rows,
	filters,
	date_label,
):
	output = BytesIO()

	workbook = Workbook()

	worksheet = workbook.active
	worksheet.title = "Daily Register"

	title_fill = PatternFill(
		fill_type="solid",
		fgColor="DDEEFF",
	)

	info_fill = PatternFill(
		fill_type="solid",
		fgColor="EEF6FC",
	)

	header_fill = PatternFill(
		fill_type="solid",
		fgColor="D8EAF8",
	)

	row_fill_1 = PatternFill(
		fill_type="solid",
		fgColor="F7FBFF",
	)

	row_fill_2 = PatternFill(
		fill_type="solid",
		fgColor="EAF2F8",
	)

	border_side = Side(
		style="thin",
		color="B9CBD9",
	)

	cell_border = Border(
		left=border_side,
		right=border_side,
		top=border_side,
		bottom=border_side,
	)

	fieldnames = []
	headers = []

	for column in columns:
		if isinstance(column, dict):
			fieldname = (
				column.get("fieldname")
				or ""
			)

			label = (
				column.get("label")
				or fieldname
				or ""
			)
		else:
			fieldname = str(column)
			label = str(column)

		fieldnames.append(fieldname)
		headers.append(label)

	total_columns = max(len(headers), 2)

	worksheet.merge_cells(
		start_row=1,
		start_column=1,
		end_row=1,
		end_column=total_columns,
	)

	title_cell = worksheet.cell(
		row=1,
		column=1,
		value=REPORT_NAME,
	)

	title_cell.font = Font(
		bold=True,
		size=16,
		color="174D7A",
	)

	title_cell.fill = title_fill

	title_cell.alignment = Alignment(
		horizontal="left",
		vertical="center",
	)

	worksheet.row_dimensions[1].height = 30

	worksheet.cell(
		row=2,
		column=1,
		value="Report Date",
	)

	worksheet.cell(
		row=2,
		column=2,
		value=date_label,
	)

	worksheet.cell(
		row=3,
		column=1,
		value="View",
	)

	worksheet.cell(
		row=3,
		column=2,
		value=(
			filters.get("report_view")
			or "Daily Consolidated Summary"
		),
	)

	for row_number in (2, 3):
		for column_number in (1, 2):
			cell = worksheet.cell(
				row=row_number,
				column=column_number,
			)

			cell.fill = info_fill
			cell.border = cell_border

			cell.alignment = Alignment(
				vertical="center",
			)

		worksheet.cell(
			row=row_number,
			column=1,
		).font = Font(
			bold=True,
			color="35566F",
		)

	header_row = 5

	for column_number, label in enumerate(
		headers,
		start=1,
	):
		cell = worksheet.cell(
			row=header_row,
			column=column_number,
			value=label,
		)

		cell.fill = header_fill

		cell.font = Font(
			bold=True,
			color="123F64",
		)

		cell.alignment = Alignment(
			horizontal="center",
			vertical="center",
			wrap_text=True,
		)

		cell.border = cell_border

	worksheet.row_dimensions[
		header_row
	].height = 32

	first_data_row = header_row + 1

	for excel_row_number, row in enumerate(
		rows,
		start=first_data_row,
	):
		row_index = (
			excel_row_number
			- first_data_row
		)

		if row_index % 2 == 0:
			row_fill = row_fill_1
		else:
			row_fill = row_fill_2

		for column_number, fieldname in enumerate(
			fieldnames,
			start=1,
		):
			value = (
				row.get(fieldname)
				if hasattr(row, "get")
				else ""
			)

			if value is None:
				value = ""

			if isinstance(
				value,
				(
					dict,
					list,
					tuple,
					set,
				),
			):
				value = str(value)

			cell = worksheet.cell(
				row=excel_row_number,
				column=column_number,
				value=value,
			)

			cell.fill = row_fill
			cell.border = cell_border

			cell.alignment = Alignment(
				vertical="top",
				wrap_text=True,
			)

	number_fields = {
		"total_actions",
		"created",
		"edited",
		"submitted",
		"cancelled",
		"workflow_changes",
		"status_changes",
		"affected_documents",
		"active_users",
		"current_draft",
		"current_submitted",
		"current_cancelled",
	}

	for column_number, fieldname in enumerate(
		fieldnames,
		start=1,
	):
		if fieldname not in number_fields:
			continue

		for row_number in range(
			first_data_row,
			worksheet.max_row + 1,
		):
			worksheet.cell(
				row=row_number,
				column=column_number,
			).number_format = "#,##0"

	if "affected_value" in fieldnames:
		amount_column = (
			fieldnames.index(
				"affected_value"
			)
			+ 1
		)

		for row_number in range(
			first_data_row,
			worksheet.max_row + 1,
		):
			worksheet.cell(
				row=row_number,
				column=amount_column,
			).number_format = "#,##0.00"

	for column_number in range(
		1,
		worksheet.max_column + 1,
	):
		max_length = 0

		for row_number in range(
			1,
			worksheet.max_row + 1,
		):
			value = worksheet.cell(
				row=row_number,
				column=column_number,
			).value

			if value is None:
				continue

			text = str(value)

			length = min(
				len(text),
				50,
			)

			if length > max_length:
				max_length = length

		width = min(
			max(
				max_length + 3,
				11,
			),
			38,
		)

		column_letter = worksheet.cell(
			row=header_row,
			column=column_number,
		).column_letter

		worksheet.column_dimensions[
			column_letter
		].width = width

	for fieldname, width in {
		"document_type": 30,
		"transaction_type": 24,
		"users": 45,
		"status_breakdown": 35,
		"latest_document": 28,
	}.items():
		if fieldname in fieldnames:
			column_number = (
				fieldnames.index(
					fieldname
				)
				+ 1
			)

			column_letter = worksheet.cell(
				row=header_row,
				column=column_number,
			).column_letter

			worksheet.column_dimensions[
				column_letter
			].width = width

	worksheet.freeze_panes = "A6"
	worksheet.sheet_view.showGridLines = False

	workbook.save(output)

	output.seek(0)

	return output.getvalue()


def format_number(value):
	try:
		return "{:,}".format(
			int(value or 0)
		)
	except Exception:
		return str(value or 0)


def metric_card(
	label,
	value,
	background="#DDEEFF",
	border="#9EC6EB",
	value_color="#185A94",
):
	return (
		"<td width='33.33%' "
		"valign='top' "
		"style='padding:5px;'>"

		"<table role='presentation' "
		"width='100%' "
		"cellpadding='0' "
		"cellspacing='0'>"

		"<tr>"

		"<td style='"
		"background:"
		+ background
		+ ";"
		"border:1px solid "
		+ border
		+ ";"
		"border-radius:8px;"
		"padding:16px 10px;"
		"text-align:center;'>"

		"<div style='"
		"font-family:Arial,sans-serif;"
		"font-size:11px;"
		"font-weight:600;"
		"color:#374151;"
		"margin-bottom:5px;'>"
		+ html.escape(str(label))
		+ "</div>"

		"<div style='"
		"font-family:Arial,sans-serif;"
		"font-size:24px;"
		"line-height:29px;"
		"font-weight:700;"
		"color:"
		+ value_color
		+ ";'>"
		+ html.escape(format_number(value))
		+ "</div>"

		"</td>"

		"</tr>"

		"</table>"

		"</td>"
	)


def build_daily_document_register_email(
	rows,
	date_label,
):
	total_actions = 0
	created = 0
	edited = 0
	submitted = 0
	cancelled = 0
	workflow_changes = 0
	status_changes = 0
	documents_touched = 0

	users = set()

	for row in rows:
		total_actions += row.get("total_actions") or 0
		created += row.get("created") or 0
		edited += row.get("edited") or 0
		submitted += row.get("submitted") or 0
		cancelled += row.get("cancelled") or 0
		workflow_changes += row.get("workflow_changes") or 0
		status_changes += row.get("status_changes") or 0
		documents_touched += row.get("affected_documents") or 0

		for user in (
			row.get("users")
			or ""
		).split(","):
			user = user.strip()

			if user:
				users.add(user)

	cards = (
		"<table role='presentation' "
		"width='100%' "
		"cellpadding='0' "
		"cellspacing='0' "
		"style='margin-bottom:18px;'>"

		"<tr>"

		+ metric_card(
			"Total Actions",
			total_actions,
			background="#D9EAFC",
			border="#8FBCE8",
			value_color="#155A96",
		)

		+ metric_card(
			"Documents Touched",
			documents_touched,
			background="#D9F1E7",
			border="#8DD0B2",
			value_color="#216E50",
		)

		+ metric_card(
			"Active Users",
			len(users),
			background="#E7DEFC",
			border="#B9A6EE",
			value_color="#613C9C",
		)

		+ "</tr>"

		"<tr>"

		+ metric_card(
			"Created",
			created,
			background="#DDF1E0",
			border="#99D0A3",
			value_color="#247137",
		)

		+ metric_card(
			"Edited",
			edited,
			background="#FDEBD5",
			border="#EAB975",
			value_color="#A45D0A",
		)

		+ metric_card(
			"Submitted",
			submitted,
			background="#D8EFEC",
			border="#8CCEC4",
			value_color="#216F62",
		)

		+ "</tr>"

		"<tr>"

		+ metric_card(
			"Cancelled",
			cancelled,
			background="#F9DDE0",
			border="#E8969E",
			value_color="#A5323D",
		)

		+ metric_card(
			"Workflow Changes",
			workflow_changes,
			background="#E5DCFA",
			border="#B6A2E9",
			value_color="#603A99",
		)

		+ metric_card(
			"Status Changes",
			status_changes,
			background="#FCEEC8",
			border="#E4C263",
			value_color="#87600C",
		)

		+ "</tr>"

		"</table>"
	)

	table_rows = ""

	for index, row in enumerate(rows):
		if index % 2 == 0:
			row_background = "#F7FBFF"
		else:
			row_background = "#EAF2F8"

		table_rows += (
			"<tr style='"
			"background:"
			+ row_background
			+ ";'>"

			"<td style='"
			"padding:7px;"
			"border:1px solid #C3D1DC;"
			"vertical-align:top;'>"
			+ html.escape(
				str(
					row.get("document_type")
					or ""
				)
			)
			+ "</td>"

			"<td style='"
			"padding:7px;"
			"border:1px solid #C3D1DC;"
			"vertical-align:top;'>"
			+ html.escape(
				str(
					row.get("transaction_type")
					or ""
				)
			)
			+ "</td>"

			"<td style='"
			"padding:7px;"
			"border:1px solid #C3D1DC;"
			"text-align:center;"
			"vertical-align:top;'>"
			+ format_number(
				row.get("total_actions")
			)
			+ "</td>"

			"<td style='"
			"padding:7px;"
			"border:1px solid #C3D1DC;"
			"text-align:center;"
			"vertical-align:top;'>"
			+ format_number(
				row.get("created")
			)
			+ "</td>"

			"<td style='"
			"padding:7px;"
			"border:1px solid #C3D1DC;"
			"text-align:center;"
			"vertical-align:top;'>"
			+ format_number(
				row.get("edited")
			)
			+ "</td>"

			"<td style='"
			"padding:7px;"
			"border:1px solid #C3D1DC;"
			"text-align:center;"
			"vertical-align:top;'>"
			+ format_number(
				row.get("submitted")
			)
			+ "</td>"

			"<td style='"
			"padding:7px;"
			"border:1px solid #C3D1DC;"
			"text-align:center;"
			"vertical-align:top;'>"
			+ format_number(
				row.get("cancelled")
			)
			+ "</td>"

			"<td style='"
			"padding:7px;"
			"border:1px solid #C3D1DC;"
			"text-align:center;"
			"vertical-align:top;'>"
			+ format_number(
				row.get("workflow_changes")
			)
			+ "</td>"

			"<td style='"
			"padding:7px;"
			"border:1px solid #C3D1DC;"
			"text-align:center;"
			"vertical-align:top;'>"
			+ format_number(
				row.get("status_changes")
			)
			+ "</td>"

			"<td style='"
			"padding:7px;"
			"border:1px solid #C3D1DC;"
			"vertical-align:top;"
			"word-break:break-word;'>"
			+ html.escape(
				str(
					row.get("users")
					or ""
				)
			)
			+ "</td>"

			"</tr>"
		)

	if not table_rows:
		table_rows = (
			"<tr>"
			"<td colspan='10' "
			"style='"
			"padding:15px;"
			"border:1px solid #C3D1DC;"
			"text-align:center;"
			"background:#F7FBFF;'>"
			"No document activity recorded for this date."
			"</td>"
			"</tr>"
		)

	return (
		"<div style='"
		"font-family:Arial,sans-serif;"
		"font-size:13px;"
		"color:#1F2937;"
		"max-width:1100px;'>"

		"<p>Dear Team,</p>"

		"<p>"
		"Please find the "
		"<b>Daily Document Register</b> "
		"for <b>"
		+ html.escape(date_label)
		+ "</b>."
		"</p>"

		+ cards +

		"<table "
		"width='100%' "
		"cellpadding='0' "
		"cellspacing='0' "
		"style='"
		"width:100%;"
		"border-collapse:collapse;"
		"font-family:Arial,sans-serif;"
		"font-size:11px;'>"

		"<thead>"

		"<tr style='"
		"background:#D8EAF8;"
		"color:#123F64;'>"

		"<th style='"
		"padding:8px 5px;"
		"border:1px solid #AFC4D4;"
		"text-align:center;'>"
		"Document Type"
		"</th>"

		"<th style='"
		"padding:8px 5px;"
		"border:1px solid #AFC4D4;"
		"text-align:center;'>"
		"Transaction Type"
		"</th>"

		"<th style='"
		"padding:8px 5px;"
		"border:1px solid #AFC4D4;"
		"text-align:center;'>"
		"Actions"
		"</th>"

		"<th style='"
		"padding:8px 5px;"
		"border:1px solid #AFC4D4;"
		"text-align:center;'>"
		"Created"
		"</th>"

		"<th style='"
		"padding:8px 5px;"
		"border:1px solid #AFC4D4;"
		"text-align:center;'>"
		"Edited"
		"</th>"

		"<th style='"
		"padding:8px 5px;"
		"border:1px solid #AFC4D4;"
		"text-align:center;'>"
		"Submitted"
		"</th>"

		"<th style='"
		"padding:8px 5px;"
		"border:1px solid #AFC4D4;"
		"text-align:center;'>"
		"Cancelled"
		"</th>"

		"<th style='"
		"padding:8px 5px;"
		"border:1px solid #AFC4D4;"
		"text-align:center;'>"
		"Workflow"
		"</th>"

		"<th style='"
		"padding:8px 5px;"
		"border:1px solid #AFC4D4;"
		"text-align:center;'>"
		"Status"
		"</th>"

		"<th style='"
		"padding:8px 5px;"
		"border:1px solid #AFC4D4;"
		"text-align:center;'>"
		"Users"
		"</th>"

		"</tr>"

		"</thead>"

		"<tbody>"
		+ table_rows
		+ "</tbody>"

		"</table>"

		"<p style='margin-top:18px;'>"
		"The complete Daily Consolidated Summary "
		"is attached in Excel format."
		"</p>"

		"<p>"
		"Regards,<br>"
		"<b>ERPNext System</b>"
		"</p>"

		"</div>"
	)